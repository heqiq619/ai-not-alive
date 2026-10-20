# -*- coding: utf-8 -*-
"""509_ch2_publish.py · 第二章 -> Word（便于导出 PDF）+ Zenodo 发布清单"""
import io
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

BASE = 'C:/Users/Administrator/Documents/deepseek-harness/default-workspace/digital-life/'
src = io.open(BASE + '论文-02-第一次测量.md', encoding='utf-8').read().split('\n')
doc = Document()
st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
in_code = False
for raw in src:
    line = raw.rstrip()
    if line.startswith('```'):
        in_code = not in_code; continue
    if in_code:
        p = doc.add_paragraph(); r = p.add_run(line)
        r.font.name = 'Consolas'; r.font.size = Pt(9)
        r.element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas'); continue
    s = line.strip()
    if not s or s == '---': continue
    if s.startswith('# ') and not s.startswith('## '):
        p = doc.add_paragraph(); r = p.add_run(s[2:].strip()); r.bold = True; r.font.size = Pt(18)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif s.startswith('## '):
        p = doc.add_paragraph(); r = p.add_run(s[3:].strip()); r.bold = True; r.font.size = Pt(13)
    elif s.startswith('### '):
        p = doc.add_paragraph(); r = p.add_run(s[4:].strip()); r.bold = True; r.font.size = Pt(11)
    else:
        doc.add_paragraph(s.replace('**', '').replace('`', ''))
doc.save(BASE + '论文-02-全文.docx')
print('OK: 论文-02-全文.docx')

checklist = '''# Zenodo 发布清单 · 第二章

## 1. 登录
https://zenodo.org → 登录（已有账号）

## 2. New upload → 填写

| 字段 | 填什么 |
|---|---|
| **Title** | 容量不能替代数据的更新 |
| **Subtitle** | 语言升维下的失配测量（scaling vs 持续学习：一个弱结论） |
| **Authors** | 何淇 — Independent Researcher |
| **Email** | 15907239031@163.com |
| **Keywords** | scaling law；持续学习；语言升维；失配测量；n-gram；cs.AI |
| **Language** | Chinese |
| **License** | CC-BY 4.0 |
| **Resource type** | Publication → Preprint |

## 3. Description（直接粘贴）

本文测量语言升维与冻结模型失配的关系，并对方法做了两次自我修正。数据：arXiv cs.AI 2018/2020/2022/2024 各 3000 篇。结果：① 词层面升维速率约 1.2%/年（6 年新词占比 7.30%）；② 等数据预算对照显示，概念新鲜度的独立价值约 0.405 bits/token，混合新旧数据的模型在新语料改善的同时旧能力不损失（13.655 vs 13.659），而只用新数据会遗忘（旧语料 13.800 vs 13.659）；③ 容量补偿实验（2~6 阶带插值 n-gram）显示旧数据与混合数据的差距稳定在 +0.86~0.94 bits/token，未见缩小，构成「容量不能替代数据更新」的弱结论。结论：语言是累积升维的，正确的更新是累积/混合而非替换；失配是数据面的问题，不是参数面的问题。局限：单领域、n-gram 级、待 transformer 验证。

## 4. Related works（关联第一章）
- Title: 现在的 AI 为什么不是生命
- Identifier: https://doi.org/10.5281/zenodo.23488543
- Relation: Is part of / Continues

## 5. 上传文件
论文-02-全文.docx（或另存 PDF）

## 6. 发布后
复制 DOI → 更新 GitHub README → 知乎发第二篇（文末挂两个 DOI）
'''
io.open(BASE + 'Zenodo发布清单-第二章.md', 'w', encoding='utf-8').write(checklist)
print('OK: Zenodo发布清单-第二章.md')
